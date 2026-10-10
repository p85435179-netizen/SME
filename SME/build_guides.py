from pathlib import Path
import re
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT

root=Path(__file__).parent
doc=Document()
sec=doc.sections[0]
sec.top_margin=sec.bottom_margin=Inches(.7)
sec.left_margin=sec.right_margin=Inches(.75)
for name in ['Normal','Title','Subtitle','Heading 1','Heading 2','Heading 3']:
    style=doc.styles[name]
    style.font.name='Calibri'
    style.font.color.rgb=RGBColor(0,0,0)
doc.styles['Normal'].font.size=Pt(11)
doc.styles['Normal'].paragraph_format.space_after=Pt(6)
doc.styles['Normal'].paragraph_format.line_spacing=1.08
for name,size in [('Title',25),('Heading 1',17),('Heading 2',14),('Heading 3',12)]:
    doc.styles[name].font.size=Pt(size)
    doc.styles[name].paragraph_format.keep_with_next=True

def inline(p,text):
    for token in re.split(r'(\*\*.*?\*\*|`[^`]+`)',text):
        r=p.add_run(token.strip('*`') if token.startswith(('**','`')) else token)
        if token.startswith('**'):r.bold=True

def heading(text):
    return re.sub(r'[^\w\s]','',text).replace('  ',' ')

def markdown(path):
    lines=path.read_text(encoding='utf-8').splitlines()
    i=0
    while i<len(lines):
        line=lines[i].strip()
        if not line:i+=1;continue
        if line.startswith('|'):
            rows=[]
            while i<len(lines) and lines[i].strip().startswith('|'):
                cells=[c.strip() for c in lines[i].strip().strip('|').split('|')]
                if not all(re.fullmatch(r'[-: ]+',c) for c in cells):rows.append(cells)
                i+=1
            table=doc.add_table(rows=0,cols=len(rows[0]))
            table.alignment=WD_TABLE_ALIGNMENT.CENTER
            table.style='Table Grid'
            for j,row in enumerate(rows):
                cells=table.add_row().cells
                for k,value in enumerate(row):
                    cells[k].vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
                    inline(cells[k].paragraphs[0],value)
                    for p in cells[k].paragraphs:
                        p.paragraph_format.space_after=Pt(5)
                        p.paragraph_format.space_before=Pt(5)
                        for r in p.runs:r.font.size=Pt(10)
                    pr=cells[k]._tc.get_or_add_tcPr()
                    shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'E7EBE6' if j==0 else ('F6F7F5' if j%2==0 else 'FFFFFF'));pr.append(shade)
                    borders=OxmlElement('w:tcBorders')
                    for edge in ['top','left','bottom','right']:
                        e=OxmlElement('w:'+edge);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');borders.append(e)
                    pr.append(borders)
                    if j==0:
                        for r in cells[k].paragraphs[0].runs:r.bold=True
                trpr=table.rows[j]._tr.get_or_add_trPr()
                if j==0:trpr.append(OxmlElement('w:tblHeader'))
                trpr.append(OxmlElement('w:cantSplit'))
            doc.add_paragraph()
            continue
        if line.startswith('#'):
            level=len(line)-len(line.lstrip('#'))
            text=heading(line.lstrip('#').strip())
            if level==1:doc.add_paragraph(text,'Title')
            else:doc.add_heading(text,level=min(level-1,3))
        elif line.startswith('- '):inline(doc.add_paragraph(style='List Bullet'),line[2:])
        elif re.match(r'^\d+\. ',line):inline(doc.add_paragraph(style='List Number'),re.sub(r'^\d+\. ','',line))
        else:inline(doc.add_paragraph(),line)
        i+=1

doc.add_paragraph('SME Financial Tracker Guides','Title')
doc.add_paragraph('System modules and role instructions','Subtitle')
doc.add_paragraph('This guide explains how the financial tracker works and how Business Owners, Accountants, and Admins use it. The first part covers system modules and criteria; the second provides role-specific workflows and access limits.')
doc.add_paragraph('Contents','Heading 1')
doc.add_paragraph('Part 1  System Module Guide')
doc.add_paragraph('Part 2  Role Guide')
doc.add_page_break()
markdown(root/'MODULE_GUIDE.md')
doc.add_page_break()
markdown(root/'ROLE_GUIDE.md')
doc.core_properties.title='SME Financial Tracker Module and Role Guides'
doc.core_properties.subject='System modules and user role workflows'
doc.save(root/'SME_Module_and_Role_Guides.docx')
print(root/'SME_Module_and_Role_Guides.docx')
